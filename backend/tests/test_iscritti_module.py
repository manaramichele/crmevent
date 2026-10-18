# Elenco Iscritti (archivio-iscritti) module: activation, editions, guided import, dedup, filters/KPI/export,
# CRUD + consent history, permissions, multi-tenant isolation, Brevo (not configured), SA endpoints, press integration.
# Run: cd /app/backend/tests && python -m pytest test_iscritti_module.py -v -n 0
import io
import os
import time
import uuid

import bcrypt
import openpyxl
import pytest
import requests
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL") or open("/app/frontend/.env").read().split("REACT_APP_BACKEND_URL=")[1].split()[0]
BASE_URL = BASE_URL.rstrip("/")
API = f"{BASE_URL}/api"
SA_EMAIL, SA_PW = "manara.michele.pro@gmail.com", "CrmEvent2026!"
TS = str(int(time.time()))
PW = "Password123!"


def _mongo():
    env = dict(l.split("=", 1) for l in open("/app/backend/.env").read().splitlines() if "=" in l)
    return MongoClient(env["MONGO_URL"].strip('"'))[env["DB_NAME"].strip('"')]


DB = _mongo()
S = {}  # shared state


def _register(tag):
    s = requests.Session()
    email = f"test_iscr{tag}_{TS}@example.com"
    r = s.post(f"{API}/auth/register-organization", json={"nome": "Test", "cognome": tag, "email": email, "password": PW,
                                                          "org_name": f"TEST_ISCR {tag} {TS}", "telefono": "+393331234567", "accept_terms": True})
    assert r.status_code in (200, 201), r.text
    me = s.get(f"{API}/auth/me").json()
    return s, email, me.get("org_id") or (me.get("organization") or {}).get("id")


def _csv(text, name="f.csv"):
    return {"file": (name, text.encode("utf-8"), "text/csv")}


def _xlsx(rows, name="f.xlsx"):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    b = io.BytesIO()
    wb.save(b)
    return {"file": (name, b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def _do_import(s, files, ed_id, gara="", mapping_override=None, new_fields=None, mode="update", tpl=None):
    up = s.post(f"{API}/iscritti/imports/upload", files=files)
    assert up.status_code == 200, up.text
    u = up.json()
    mapping = dict(u["suggested"])
    if mapping_override:
        mapping.update(mapping_override)
    pv = s.post(f"{API}/iscritti/imports/{u['id']}/preview", json={"edizione_id": ed_id, "gara": gara, "mapping": mapping, "new_fields": new_fields or []})
    assert pv.status_code == 200, pv.text
    rn = s.post(f"{API}/iscritti/imports/{u['id']}/run", json={"mode": mode, "template_name": tpl})
    assert rn.status_code == 200, rn.text
    for _ in range(40):
        g = s.get(f"{API}/iscritti/imports/{u['id']}").json()
        if g["status"] in ("completato", "errore"):
            break
        time.sleep(0.5)
    return u, pv.json(), g


@pytest.fixture(scope="module", autouse=True)
def setup_and_cleanup():
    sa = requests.Session()
    assert sa.post(f"{API}/auth/login", json={"email": SA_EMAIL, "password": SA_PW}).status_code == 200
    S["sa"] = sa
    S["a"], S["a_email"], S["a_org"] = _register("A")
    S["b"], S["b_email"], S["b_org"] = _register("B")
    # non-admin member of org A inserted directly (role 'user')
    uid = f"user_test_{uuid.uuid4().hex[:10]}"
    S["m_email"] = f"test_iscrm_{TS}@example.com"
    now = "2026-01-01T00:00:00+00:00"
    DB.users.insert_one({"user_id": uid, "email": S["m_email"], "name": "Test Member", "password_hash": bcrypt.hashpw(PW.encode(), bcrypt.gensalt()).decode(),
                         "role": "member", "auth_provider": "password", "telefono": "+393331112222", "org_id": S["a_org"], "picture": "", "active": True,
                         "accepted_terms_at": now, "terms_version": "x", "privacy_version": "x", "created_at": now, "registered_at": now})
    DB.memberships.insert_one({"id": uuid.uuid4().hex, "user_id": uid, "org_id": S["a_org"], "role": "user", "active": True, "permissions": {}, "created_at": now, "updated_at": now})
    m = requests.Session()
    lr = m.post(f"{API}/auth/login", json={"email": S["m_email"], "password": PW})
    S["m"] = m if lr.status_code == 200 else None
    S["m_login"] = lr.status_code
    yield
    orgs = [S["a_org"], S["b_org"]]
    users = [u["user_id"] for u in DB.users.find({"$or": [{"org_id": {"$in": orgs}}, {"email": {"$in": [S["a_email"], S["b_email"], S["m_email"]]}}]})]
    for c in DB.list_collection_names():
        if c != "audit_logs":
            DB[c].delete_many({"org_id": {"$in": orgs}})
    DB.memberships.delete_many({"user_id": {"$in": users}})
    DB.users.delete_many({"user_id": {"$in": users}})
    DB.organizations.delete_many({"id": {"$in": orgs}})


class TestIscritti:
    def test_01_inactive_403_and_status(self):
        a = S["a"]
        st = a.get(f"{API}/iscritti/status")
        assert st.status_code == 200 and st.json()["active"] is False
        assert a.get(f"{API}/iscritti").status_code == 403
        assert a.get(f"{API}/iscritti/kpi").status_code == 403
        assert a.get(f"{API}/iscritti/edizioni").status_code == 403

    def test_02_sa_activate_and_marketplace(self):
        for o in (S["a_org"], S["b_org"]):
            r = S["sa"].post(f"{API}/platform/iscritti/orgs/{o}/activate", json={"active": True})
            assert r.status_code == 200, r.text
        assert S["a"].get(f"{API}/iscritti/status").json()["active"] is True
        assert S["a"].get(f"{API}/iscritti").status_code == 200
        svcs = S["a"].get(f"{API}/marketplace/services")
        assert svcs.status_code == 200, svcs.text
        data = svcs.json()
        items = data if isinstance(data, list) else data.get("services") or data.get("items") or []
        hit = [x for x in items if x.get("key") == "archivio-iscritti"]
        assert hit and hit[0]["name"] == "Elenco Iscritti", hit
        print("marketplace item:", {k: hit[0].get(k) for k in ("key", "name", "active", "is_active", "status", "org_active")})
        assert hit[0].get("active") or hit[0].get("is_active") or hit[0].get("org_active") or hit[0].get("activated")

    def test_03_events_and_editions(self):
        a = S["a"]
        ev = a.post(f"{API}/events", json={"nome": f"TEST_Maratona {TS}", "data_inizio": "2026-05-10"})
        assert ev.status_code in (200, 201), ev.text
        S["ev"] = ev.json()["id"]
        e1 = a.post(f"{API}/iscritti/edizioni", json={"event_id": S["ev"], "anno": 2026, "edizione": "10a", "gare": ["42K", "21K"]})
        e2 = a.post(f"{API}/iscritti/edizioni", json={"event_id": S["ev"], "anno": 2019, "edizione": "3a", "gare": ["42K"]})
        assert e1.status_code == 200 and e2.status_code == 200, e1.text + e2.text
        S["ed26"], S["ed19"] = e1.json()["id"], e2.json()["id"]
        dup = a.post(f"{API}/iscritti/edizioni", json={"event_id": S["ev"], "anno": 2026, "edizione": "10a", "gare": []})
        assert dup.status_code == 409
        lst = a.get(f"{API}/iscritti/edizioni").json()
        assert {S["ed26"], S["ed19"]} <= {e["id"] for e in lst["edizioni"]}

    def test_04_import_english_semicolon_csv(self):
        csv_txt = ("First Name;Last Name;E-mail;Mobile;Club;Race;Registration Date;Newsletter;Shirt Size\n"
                   "Mario;Rossi;mario.TEST@example.com;333111;ASD Run;42K;2026-01-05;si;L\n"
                   "Anna;Bianchi;anna.test@example.com;333222;Team X;21K;05/01/2026;no;M\n"
                   "Luca;Verdi;luca.test@example.com;;;42K;;;S\n"
                   "Mario;Rossi;mario.test@example.com;333111;ASD Run;42K;2026-01-05;si;L\n"
                   ";;bad-email;;;;;;\n")
        u, pv, g = _do_import(S["a"], _csv(csv_txt, "en.csv"), S["ed26"], new_fields=["Shirt Size"], tpl="TEST_tpl_en")
        sug = u["suggested"]
        print("suggested:", sug)
        assert sug["First Name"] == "nome" and sug["Last Name"] == "cognome" and sug["E-mail"] == "email"
        assert sug["Mobile"] == "cellulare" and sug["Club"] == "societa" and sug["Race"] == "gara"
        assert sug["Registration Date"] == "data_iscrizione" and sug["Newsletter"] == "consenso_marketing"
        st = pv["stats"]
        print("preview:", pv)
        assert st["total"] == 5 and st["new"] == 3 and st["duplicates"] == 1 and st["errors"] == 1
        assert g["status"] == "completato", g
        assert g["result"]["created"] == 3 and g["result"]["participations"] == 3
        assert g.get("errors")
        items = S["a"].get(f"{API}/iscritti", params={"limit": 200}).json()["items"]
        by = {i["email"]: i for i in items}
        assert by["mario.test@example.com"]["consenso_marketing"] is True and by["mario.test@example.com"]["consenso_fonte"]
        assert by["anna.test@example.com"]["consenso_marketing"] is False
        assert by["luca.test@example.com"]["consenso_marketing"] is False
        assert by["mario.test@example.com"]["custom"].get("shirt_size") == "L"
        S["mario"] = by["mario.test@example.com"]["id"]
        st2 = S["a"].get(f"{API}/iscritti/status").json()
        assert any(c["key"] == "shirt_size" for c in st2["custom_fields"])

    def test_05_reimport_same_no_dupes(self):
        csv_txt = "First Name;Last Name;E-mail;Race;Newsletter\nMario;Rossi;mario.test@example.com;42K;si\nAnna;Bianchi;anna.test@example.com;21K;no\n"
        u, pv, g = _do_import(S["a"], _csv(csv_txt, "en2.csv"), S["ed26"])
        assert pv["stats"]["new"] == 0 and pv["stats"]["existing"] == 2 and pv["stats"]["already_in_edition"] == 2
        assert g["result"]["created"] == 0 and g["result"]["participations"] == 0 and g["result"]["skipped_participations"] == 2
        assert S["a"].get(f"{API}/iscritti").json()["total"] == 3

    def test_06_italian_comma_csv_other_edition_and_remap(self):
        csv_txt = "Nome,Cognome,Email,Gara,Telefono,Commenti\nMario,Rossi,mario.test@example.com,42K,999,nota1\nGiulia,Neri,giulia.test@example.com,42K,888,nota2\n"
        # manual remap: Telefono -> ignore, Commenti -> societa
        u, pv, g = _do_import(S["a"], _csv(csv_txt, "it.csv"), S["ed19"], mapping_override={"Telefono": "", "Commenti": "societa"})
        assert u["suggested"]["Nome"] == "nome" and u["suggested"]["Gara"] == "gara"
        assert g["status"] == "completato" and g["result"]["created"] == 1 and g["result"]["participations"] == 2
        det = S["a"].get(f"{API}/iscritti/{S['mario']}").json()
        assert len(det["partecipazioni"]) == 2
        items = {i["email"]: i for i in S["a"].get(f"{API}/iscritti", params={"q": "giulia"}).json()["items"]}
        gi = items["giulia.test@example.com"]
        assert gi["societa"] == "nota2" and not gi.get("cellulare") and gi["consenso_marketing"] is False
        S["giulia"] = gi["id"]

    def test_07_xlsx_import(self):
        rows = [["Name", "Surname", "Email Address", "Race", "Newsletter"], ["Paolo", "Gialli", "paolo.test@example.com", "21K", "yes"]]
        u, pv, g = _do_import(S["a"], _xlsx(rows), S["ed26"])
        assert u["suggested"]["Name"] == "nome" and u["suggested"]["Surname"] == "cognome" and u["suggested"]["Email Address"] == "email"
        assert g["status"] == "completato" and g["result"]["created"] == 1
        assert DB.iscritti_import_files.count_documents({"import_id": u["id"]}) == 0

    def test_08_template_reused(self):
        csv_txt = "First Name;Last Name;E-mail;Mobile;Club;Race;Registration Date;Newsletter;Shirt Size\nZed;Test;zed.test@example.com;;;42K;;no;XL\n"
        up = S["a"].post(f"{API}/iscritti/imports/upload", files=_csv(csv_txt, "tpl.csv")).json()
        assert up.get("template") == "TEST_tpl_en", up.get("template")
        assert up["suggested"]["Shirt Size"] == "custom:shirt_size"

    def test_09_filters_kpi_export(self):
        a = S["a"]
        t = a.get(f"{API}/iscritti").json()["total"]
        assert t == 5
        assert a.get(f"{API}/iscritti", params={"edizione_id": S["ed19"]}).json()["total"] == 2
        assert a.get(f"{API}/iscritti", params={"anno": "2026", "gara": "21K"}).json()["total"] == 2
        assert a.get(f"{API}/iscritti", params={"event_id": S["ev"], "consenso": "si"}).json()["total"] == 2
        assert a.get(f"{API}/iscritti", params={"stato": "confermata", "q": "rossi"}).json()["total"] == 1
        srt = a.get(f"{API}/iscritti", params={"sort": "email", "dir": -1, "limit": 2, "page": 1}).json()
        assert len(srt["items"]) == 2 and srt["items"][0]["email"] >= srt["items"][1]["email"]
        p3 = a.get(f"{API}/iscritti", params={"limit": 2, "page": 3}).json()
        assert len(p3["items"]) == 1
        k = a.get(f"{API}/iscritti/kpi").json()
        print("kpi:", k)
        assert k["total"] == 5 and k["recurring"] == 1 and k["new_30d"] == 5 and k["with_consent"] == 2
        assert k["by_event"][0]["n"] == 5 and len(k["by_edition"]) == 2
        c = a.get(f"{API}/iscritti/export", params={"format": "csv"})
        assert c.status_code == 200 and "mario.test@example.com" in c.text and "Shirt Size" in c.text
        x = a.get(f"{API}/iscritti/export", params={"format": "xlsx", "consenso": "si"})
        assert x.status_code == 200 and x.content[:2] == b"PK"
        ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
        assert ws.max_row == 3

    def test_10_manual_crud_and_consent_history(self):
        a = S["a"]
        r = a.post(f"{API}/iscritti", json={"nome": "Man", "cognome": "Ual", "email": "manual.test@example.com",
                                            "partecipazione": {"edizione_id": S["ed26"], "gara": "21K"}})
        assert r.status_code == 200, r.text
        pid = r.json()["id"]
        assert a.post(f"{API}/iscritti", json={"nome": "X", "cognome": "Y", "email": "MANUAL.test@example.com"}).status_code == 409
        up = a.put(f"{API}/iscritti/{pid}", json={"nome": "Manu", "cognome": "Ual", "email": "manual.test@example.com", "consenso_marketing": True, "consenso_fonte": "Modulo carta"})
        assert up.status_code == 200, up.text
        d = up.json()
        assert d["nome"] == "Manu" and d["consenso_marketing"] is True and len(d["consent_history"]) == 1
        up2 = a.put(f"{API}/iscritti/{pid}", json={"nome": "Manu", "cognome": "Ual", "email": "manual.test@example.com", "consenso_marketing": False})
        assert len(up2.json()["consent_history"]) == 2 and up2.json()["consent_history"][1]["value"] is False
        assert len(d["partecipazioni"]) == 1
        S["manual"] = pid
        # last_year should be set for manually added participation (retention cron relies on it)
        doc = DB.iscritti.find_one({"id": pid})
        S["manual_last_year"] = doc.get("last_year")
        print("manual last_year:", doc.get("last_year"))

    def test_11_anonymize_and_delete(self):
        a = S["a"]
        assert a.post(f"{API}/iscritti/{S['giulia']}/anonymize").status_code == 200
        g = a.get(f"{API}/iscritti/{S['giulia']}").json()
        assert g["anonymized"] is True and g["email"] == "" and g["nome"] == ""
        assert a.get(f"{API}/iscritti", params={"q": "giulia"}).json()["total"] == 0
        assert a.delete(f"{API}/iscritti/{S['manual']}").status_code == 200
        assert a.get(f"{API}/iscritti/{S['manual']}").status_code == 404
        assert DB.iscritti_partecipazioni.count_documents({"iscritto_id": S["manual"]}) == 0

    def test_12_member_permissions(self):
        assert S["m"] is not None, f"member login failed {S['m_login']}"
        m = S["m"]
        st = m.get(f"{API}/iscritti/status").json()
        assert st["active"] is True and st["can_edit"] is False
        assert m.get(f"{API}/iscritti").status_code == 200
        assert m.get(f"{API}/iscritti/{S['mario']}").status_code == 200
        assert m.post(f"{API}/iscritti", json={"nome": "No", "cognome": "Way", "email": "noway.test@example.com"}).status_code == 403
        assert m.put(f"{API}/iscritti/{S['mario']}", json={"nome": "H", "cognome": "X", "email": "mario.test@example.com"}).status_code == 403
        assert m.delete(f"{API}/iscritti/{S['mario']}").status_code == 403
        assert m.get(f"{API}/iscritti/export").status_code == 403
        assert m.post(f"{API}/iscritti/imports/upload", files=_csv("Nome,Cognome\nA,B\n")).status_code == 403
        assert m.post(f"{API}/iscritti/edizioni", json={"event_id": S["ev"], "anno": 2027}).status_code == 403

    def test_13_isolation(self):
        b = S["b"]
        assert b.get(f"{API}/iscritti").json()["total"] == 0
        assert b.get(f"{API}/iscritti/{S['mario']}").status_code == 404
        assert b.put(f"{API}/iscritti/{S['mario']}", json={"nome": "H", "cognome": "X", "email": "h@x.com"}).status_code == 404
        assert b.delete(f"{API}/iscritti/{S['mario']}").status_code == 404
        assert b.post(f"{API}/iscritti/{S['mario']}/anonymize").status_code == 404
        assert b.post(f"{API}/iscritti/edizioni", json={"event_id": S["ev"], "anno": 2026}).status_code == 404
        assert b.get(f"{API}/iscritti/kpi").json()["total"] == 0
        # same email in org B is allowed (per-org dedup)
        assert b.post(f"{API}/iscritti", json={"nome": "Mario", "cognome": "Rossi", "email": "mario.test@example.com"}).status_code == 200

    def test_14_brevo_not_configured_and_sa(self):
        r = S["a"].post(f"{API}/iscritti/brevo/sync")
        assert r.status_code == 200 and r.json()["status"] == "non_configurato"
        sa = S["sa"]
        o = sa.get(f"{API}/platform/iscritti/orgs")
        assert o.status_code == 200
        row = next(x for x in o.json()["orgs"] if x["org_id"] == S["a_org"])
        assert row["active"] and row["iscritti"] == 4 and row["via"] == "super_admin"
        assert sa.post(f"{API}/platform/iscritti/orgs/{S['a_org']}/sync").status_code == 200
        f = sa.post(f"{API}/platform/iscritti/orgs/{S['a_org']}/force-sync")
        assert f.status_code == 200, f.text
        acts = {x["action"] for x in DB.audit_logs.find({"org_id": S["a_org"]}, {"action": 1})} if "audit_logs" in DB.list_collection_names() else set()
        if not acts:
            acts = {x.get("action") for c in DB.list_collection_names() if "audit" in c for x in DB[c].find({"org_id": S["a_org"]})}
        print("audit actions:", acts)
        assert {"iscritti_brevo_resync", "iscritti_brevo_force_sync", "iscritti_service_activated"} <= acts
        assert S["a"].post(f"{API}/platform/iscritti/orgs/{S['a_org']}/sync").status_code in (401, 403)

    def test_15_press_recipients_consent_only(self):
        r = S["sa"].post(f"{API}/platform/press/orgs/{S['a_org']}/activate", json={"active": True})
        assert r.status_code == 200, r.text
        rid = uuid.uuid4().hex
        DB.press_releases.insert_one({"id": rid, "org_id": S["a_org"], "event_id": S["ev"], "tipo": "pre", "titolo": "TEST", "testo": "x", "status": "bozza",
                                      "citazioni": [], "foto": [], "created_at": "2026-01-01T00:00:00+00:00", "updated_at": "2026-01-01T00:00:00+00:00"})
        rc = S["a"].post(f"{API}/press/releases/{rid}/recipients", json={"iscritti": {"edizione_id": S["ed26"]}})
        assert rc.status_code == 200, rc.text
        print("press recipients:", rc.json())
        # ed26 consented persons: mario (si), paolo (yes); anna no, luca missing
        assert rc.json()["iscritti"] == 2 and rc.json()["iscritti_available"] is True
        rc19 = S["a"].post(f"{API}/press/releases/{rid}/recipients", json={"iscritti": {"edizione_id": S["ed19"]}})
        assert rc19.json()["iscritti"] == 1

    def test_16_deactivate_blocks(self):
        S["sa"].post(f"{API}/platform/iscritti/orgs/{S['b_org']}/activate", json={"active": False})
        assert S["b"].get(f"{API}/iscritti").status_code == 403
        assert S["b"].get(f"{API}/iscritti/status").json()["active"] is False
