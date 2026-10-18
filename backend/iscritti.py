"""Elenco Iscritti (Marketplace "archivio-iscritti"): anagrafica unica per organizzazione, partecipazioni per edizione,
importazione guidata XLSX/XLS/CSV asincrona, sincronizzazione Brevo (solo consenso marketing), servizio centrale per gli altri moduli."""
import asyncio
import csv
import io
import logging
import os
import re
import secrets
import unicodedata
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, UploadFile
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel, Field

import brevo_client
import brevo_org_lists

log = logging.getLogger("iscritti")
SERVICE_KEY = "archivio-iscritti"
EMAIL_RX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
UPLOAD_CHUNK = 4 * 1024 * 1024
MAX_UPLOAD = 25 * 1024 * 1024
RETENTION_YEARS = 5
PERSON_FIELDS = ["nome", "cognome", "email", "cellulare", "data_nascita", "sesso", "nazionalita", "societa", "codice_fiscale",
                 "numero_tessera", "federazione", "consenso_marketing", "consenso_fonte", "consenso_data", "note"]
PART_FIELDS = ["tipologia", "gara", "stato", "data_iscrizione"]
LABELS = {"nome": "Nome", "cognome": "Cognome", "email": "Email", "cellulare": "Cellulare", "data_nascita": "Data di nascita", "sesso": "Sesso",
          "nazionalita": "Nazionalità", "societa": "Società sportiva", "codice_fiscale": "Codice fiscale", "numero_tessera": "Numero tessera",
          "federazione": "Federazione o ente", "tipologia": "Tipologia iscrizione", "gara": "Distanza o gara", "stato": "Stato iscrizione",
          "data_iscrizione": "Data iscrizione", "consenso_marketing": "Consenso marketing", "consenso_fonte": "Fonte del consenso",
          "consenso_data": "Data del consenso", "note": "Note"}
SYN = {"nome": ["nome", "first name", "firstname", "name", "given name", "prenom"], "cognome": ["cognome", "last name", "lastname", "surname", "family name"],
       "email": ["email", "e-mail", "mail", "indirizzo email", "email address"], "cellulare": ["cellulare", "mobile", "telefono", "phone", "cell", "tel"],
       "data_nascita": ["data di nascita", "data nascita", "birth date", "birthdate", "date of birth", "dob", "nascita"],
       "sesso": ["sesso", "gender", "sex", "genere"], "nazionalita": ["nazionalita", "nationality", "nazione", "country", "paese"],
       "societa": ["societa", "societa sportiva", "club", "team", "squadra", "asd"], "codice_fiscale": ["codice fiscale", "cf", "fiscal code", "tax code"],
       "numero_tessera": ["tessera", "numero tessera", "n tessera", "card", "license", "licence", "membership"],
       "federazione": ["federazione", "ente", "federation", "eps"], "tipologia": ["tipologia", "tipologia iscrizione", "tipo iscrizione", "category", "categoria", "type"],
       "gara": ["gara", "distanza", "race", "distance", "percorso"], "stato": ["stato", "stato iscrizione", "status"],
       "data_iscrizione": ["data iscrizione", "registration date", "data registrazione", "iscritto il", "signup date"],
       "consenso_marketing": ["consenso marketing", "marketing", "newsletter", "consenso", "privacy marketing", "optin", "opt in"],
       "consenso_fonte": ["fonte consenso"], "consenso_data": ["data consenso"], "note": ["note", "notes", "commenti"]}
STATI = ["confermata", "in_attesa", "annullata", "ritirato"]
TRUE = {"si", "sì", "s", "yes", "y", "true", "1", "x", "ok", "vero"}


def _now():
    return datetime.now(timezone.utc).isoformat()


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return "" if v != v else (str(int(v)) if v.is_integer() else str(v))
    return str(v).strip()


def _date(v: str) -> Optional[str]:
    v = (v or "").strip()[:19]
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(v, f).date().isoformat()
        except ValueError:
            continue
    return None


def _dedup_key(p: dict) -> Optional[str]:
    if p.get("email"):
        return "e:" + p["email"]
    if p.get("codice_fiscale"):
        return "cf:" + p["codice_fiscale"].upper()
    if p.get("nome") and p.get("cognome") and p.get("data_nascita"):
        return "n:" + _norm(f"{p['nome']} {p['cognome']}") + ":" + p["data_nascita"]
    return None


def suggest(columns: List[str], custom: List[dict]) -> Dict[str, str]:
    out, used = {}, set()
    for col in columns:
        n = _norm(col)
        hit = next((f for f, syns in SYN.items() if f not in used and n in syns), None)
        hit = hit or next((f for f, syns in SYN.items() if f not in used and any(len(s) > 3 and s in n for s in syns)), None)
        hit = hit or next((f"custom:{c['key']}" for c in custom if _norm(c["label"]) == n), None)
        out[col] = hit or ""
        if hit:
            used.add(hit)
    return out


def read_rows(data: bytes, ext: str):
    """Ritorna (colonne, righe come liste di stringhe)."""
    if ext == "csv":
        raw = data
        text = next((raw.decode(enc) for enc in ("utf-8-sig", "cp1252", "latin-1") if _try(raw, enc)), raw.decode("latin-1"))
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|") if text.strip() else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    elif ext == "xlsx":
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [[_cell(c) for c in r] for r in ws.iter_rows(values_only=True) if any(_cell(c) for c in r)]
    else:
        import xlrd
        sh = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        rows = []
        for i in range(sh.nrows):
            r = []
            for c in sh.row(i):
                r.append(xlrd.xldate_as_datetime(c.value, 0).strftime("%Y-%m-%d") if c.ctype == 3 else _cell(c.value))
            if any(r):
                rows.append(r)
    if not rows:
        return [], []
    cols, seen = [], {}
    for i, c in enumerate(rows[0]):
        c = (c or "").strip() or f"Colonna {i + 1}"
        seen[c] = seen.get(c, 0) + 1
        cols.append(c if seen[c] == 1 else f"{c} ({seen[c]})")
    return cols, [[(r[i] if i < len(r) else "").strip() for i in range(len(cols))] for r in rows[1:]]


def _try(raw, enc):
    try:
        raw.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def parse_row(cols: list, row: list, mapping: dict) -> tuple:
    """(persona, partecipazione, custom, errori)."""
    p, part, custom, errs = {}, {}, {}, []
    for col, val in zip(cols, row):
        f = mapping.get(col)
        if not f or not val:
            continue
        if f.startswith("custom:"):
            custom[f[7:]] = val[:500]
        elif f in PART_FIELDS:
            part[f] = val[:200]
        elif f in PERSON_FIELDS:
            p[f] = val[:500]
    if "email" in p:
        p["email"] = p["email"].strip().lower()
        if not EMAIL_RX.match(p["email"]):
            errs.append(f"email non valida ({p['email'][:60]})")
            p.pop("email")
    for k in ("data_nascita", "consenso_data"):
        if k in p:
            d = _date(p[k])
            if d:
                p[k] = d
            else:
                errs.append(f"{LABELS[k].lower()} non valida")
                p.pop(k)
    if "data_iscrizione" in part:
        part["data_iscrizione"] = _date(part["data_iscrizione"]) or part.pop("data_iscrizione") and None
    if "consenso_marketing" in p:
        p["consenso_marketing"] = _norm(p["consenso_marketing"]) in TRUE
    if "codice_fiscale" in p:
        p["codice_fiscale"] = p["codice_fiscale"].upper().replace(" ", "")
    if not (p.get("email") or (p.get("nome") and p.get("cognome"))):
        errs.append("servono almeno email oppure nome e cognome")
    elif not _dedup_key(p):
        errs.append("impossibile identificare la persona: servono email, codice fiscale o nome+cognome+data di nascita")
    return p, part, custom, errs


class EdizioneIn(BaseModel):
    event_id: str
    anno: int = Field(ge=1950, le=2100)
    edizione: str = Field(min_length=1, max_length=40)
    gare: List[str] = []


class FieldsIn(BaseModel):
    custom_fields: List[Dict[str, str]]


class PreviewIn(BaseModel):
    edizione_id: str
    gara: str = ""
    mapping: Dict[str, str]
    new_fields: List[str] = []


class RunIn(BaseModel):
    mode: str = Field(pattern="^(update|skip)$")
    template_name: Optional[str] = Field(default=None, max_length=80)


class IscrittoIn(BaseModel):
    nome: str = Field(default="", max_length=120)
    cognome: str = Field(default="", max_length=120)
    email: str = Field(default="", max_length=200)
    cellulare: str = ""
    data_nascita: str = ""
    sesso: str = ""
    nazionalita: str = ""
    societa: str = ""
    codice_fiscale: str = ""
    numero_tessera: str = ""
    federazione: str = ""
    note: str = Field(default="", max_length=2000)
    consenso_marketing: bool = False
    consenso_fonte: str = Field(default="", max_length=200)
    custom: Dict[str, str] = {}
    partecipazione: Optional[Dict[str, str]] = None


class SegmentIn(BaseModel):
    filters: Dict[str, str] = {}


class ActivateIn(BaseModel):
    active: bool


def build(db, deps: dict):
    r = APIRouter(prefix="/api")
    require_member, sa, record_audit = deps["require_member"], deps["require_superadmin"], deps["record_audit"]
    app_url = deps["app_url"]

    async def is_active(org_id: str) -> bool:
        if await db.org_services.find_one({"org_id": org_id, "service": SERVICE_KEY, "active": True}, {"_id": 1}):
            return True
        svc = await db.marketplace_services.find_one({"key": SERVICE_KEY}, {"_id": 0, "id": 1})
        return bool(svc and await db.marketplace_purchases.find_one({"org_id": org_id, "service_id": svc["id"], "status": {"$in": ["active", "past_due"]}}, {"_id": 1}))

    async def member(user: dict = Depends(require_member)) -> dict:
        if not await is_active(user["org_id"]):
            raise HTTPException(status_code=403, detail="Il modulo Elenco Iscritti non è attivo per questa organizzazione")
        return user

    async def editor(user: dict = Depends(member)) -> dict:
        if user.get("org_role") not in ("admin_org", "superadmin"):
            raise HTTPException(status_code=403, detail="Operazione riservata agli amministratori dell'organizzazione")
        return user

    async def settings(org_id: str) -> dict:
        return await db.iscritti_settings.find_one({"org_id": org_id}, {"_id": 0}) or {"org_id": org_id, "custom_fields": []}

    def _pub(p: dict) -> dict:
        return {k: v for k, v in p.items() if k not in ("_id", "unsub_token", "dedup_key")}

    # ---------------- servizio centrale (usato anche da Comunicati stampa e futuri moduli) ----------------
    async def participation_query(org_id: str, f: dict) -> Optional[dict]:
        q = {"org_id": org_id}
        if f.get("edizione_id"):
            q["edizione_id"] = f["edizione_id"]
        if f.get("event_id"):
            q["event_id"] = f["event_id"]
        if f.get("anno"):
            q["anno"] = int(f["anno"])
        if f.get("gara"):
            q["gara"] = f["gara"]
        if f.get("stato"):
            q["stato"] = f["stato"]
        return q if len(q) > 1 else None

    async def person_query(org_id: str, f: dict) -> dict:
        q = {"org_id": org_id, "anonymized": {"$ne": True}}
        pq = await participation_query(org_id, f)
        if pq:
            q["id"] = {"$in": await db.iscritti_partecipazioni.distinct("iscritto_id", pq)}
        if f.get("q"):
            rx = {"$regex": re.escape(f["q"].strip()[:80]), "$options": "i"}
            q["$or"] = [{"nome": rx}, {"cognome": rx}, {"email": rx}, {"societa": rx}]
        if f.get("consenso") in ("si", "no"):
            q["consenso_marketing"] = f["consenso"] == "si"
        return q

    async def select_recipients(org_id: str, filters: dict, require_consent: bool = True) -> list:
        """Iscritti con email valida per i filtri dati; di default solo con consenso marketing e non disiscritti."""
        q = await person_query(org_id, filters)
        q["email"] = {"$nin": [None, ""]}
        if require_consent:
            q["consenso_marketing"] = True
            q["unsubscribed"] = {"$ne": True}
        return await db.iscritti.find(q, {"_id": 0, "id": 1, "email": 1, "nome": 1, "cognome": 1, "unsub_token": 1}).to_list(100000)

    def unsub_url(token: str) -> str:
        return f"{app_url.rstrip('/')}/api/public/iscritti/unsubscribe/{token}"

    # ---------------- stato, campi, edizioni ----------------
    @r.get("/iscritti/status")
    async def status(user: dict = Depends(require_member)):
        act = await is_active(user["org_id"])
        st = await settings(user["org_id"]) if act else {}
        return {"active": act, "can_edit": user.get("org_role") in ("admin_org", "superadmin"), "labels": LABELS, "stati": STATI,
                "custom_fields": st.get("custom_fields", []), "brevo": {k: st.get(k) for k in ("brevo_list_id", "brevo_list_name", "brevo_last_sync")}}

    @r.put("/iscritti/fields")
    async def put_fields(body: FieldsIn, user: dict = Depends(editor)):
        out, keys = [], set()
        for c in body.custom_fields[:40]:
            label = (c.get("label") or "").strip()[:60]
            key = c.get("key") or re.sub(r"[^a-z0-9]+", "_", _norm(label)).strip("_")[:40]
            if label and key and key not in keys:
                keys.add(key)
                out.append({"key": key, "label": label})
        await db.iscritti_settings.update_one({"org_id": user["org_id"]}, {"$set": {"custom_fields": out, "updated_at": _now()}}, upsert=True)
        return {"custom_fields": out}

    @r.get("/iscritti/edizioni")
    async def list_edizioni(user: dict = Depends(member)):
        rows = await db.iscritti_edizioni.find({"org_id": user["org_id"]}, {"_id": 0}).sort([("anno", -1)]).to_list(2000)
        counts = {c["_id"]: c["n"] for c in await db.iscritti_partecipazioni.aggregate([{"$match": {"org_id": user["org_id"]}},
                                                                                         {"$group": {"_id": "$edizione_id", "n": {"$sum": 1}}}]).to_list(5000)}
        evs = {e["id"]: e for e in await db.events.find({"org_id": user["org_id"]}, {"_id": 0, "id": 1, "nome": 1, "edizione": 1}).to_list(5000)}
        return {"edizioni": [{**e, "event_nome": (evs.get(e["event_id"]) or {}).get("nome") or e.get("event_nome"), "iscritti": counts.get(e["id"], 0)} for e in rows],
                "events": [{"id": e["id"], "nome": e.get("nome")} for e in evs.values()]}

    @r.post("/iscritti/edizioni")
    async def create_edizione(body: EdizioneIn, user: dict = Depends(editor)):
        ev = await db.events.find_one({"id": body.event_id, "org_id": user["org_id"]}, {"_id": 0, "nome": 1})
        if not ev:
            raise HTTPException(status_code=404, detail="Evento non trovato")
        if await db.iscritti_edizioni.find_one({"org_id": user["org_id"], "event_id": body.event_id, "anno": body.anno, "edizione": body.edizione.strip()}):
            raise HTTPException(status_code=409, detail="Questa edizione esiste già")
        doc = {"id": uuid.uuid4().hex, "org_id": user["org_id"], "event_id": body.event_id, "event_nome": ev["nome"], "anno": body.anno,
               "edizione": body.edizione.strip(), "gare": [g.strip()[:80] for g in body.gare if g.strip()][:30], "created_at": _now(), "created_by": user["user_id"]}
        await db.iscritti_edizioni.insert_one(dict(doc))
        return doc

    @r.delete("/iscritti/edizioni/{eid}")
    async def delete_edizione(eid: str, user: dict = Depends(editor)):
        if await db.iscritti_partecipazioni.count_documents({"org_id": user["org_id"], "edizione_id": eid}):
            raise HTTPException(status_code=400, detail="L'edizione ha iscritti collegati: non può essere eliminata")
        await db.iscritti_edizioni.delete_one({"id": eid, "org_id": user["org_id"]})
        return {"ok": True}

    async def _edizione(org_id: str, eid: str) -> dict:
        e = await db.iscritti_edizioni.find_one({"id": eid, "org_id": org_id}, {"_id": 0})
        if not e:
            raise HTTPException(status_code=404, detail="Edizione non trovata")
        return e

    # ---------------- importazione ----------------
    @r.post("/iscritti/imports/upload")
    async def upload(file: UploadFile = File(...), user: dict = Depends(editor)):
        ext = (file.filename or "").rsplit(".", 1)[-1].lower()
        if ext not in ("xlsx", "xls", "csv"):
            raise HTTPException(status_code=400, detail="Formato non supportato: carica un file XLSX, XLS o CSV")
        data = await file.read(MAX_UPLOAD + 1)
        if len(data) > MAX_UPLOAD:
            raise HTTPException(status_code=413, detail="File troppo grande (massimo 25 MB)")
        iid = uuid.uuid4().hex
        try:
            cols, rows = await asyncio.to_thread(read_rows, data, ext)
        except Exception as e:  # noqa: BLE001
            log.warning("iscritti read failed: %s", type(e).__name__)
            raise HTTPException(status_code=400, detail="Impossibile leggere il file: verifica che non sia protetto o danneggiato")
        if not cols or not rows:
            raise HTTPException(status_code=400, detail="Il file non contiene righe da importare")
        # File temporaneo nel database (a blocchi), cancellato a fine importazione: nessun file sul server.
        for n, i in enumerate(range(0, len(data), UPLOAD_CHUNK)):
            await db.iscritti_import_files.insert_one({"import_id": iid, "org_id": user["org_id"], "n": n, "data": data[i:i + UPLOAD_CHUNK], "created_at": _now()})
        st = await settings(user["org_id"])
        templates = await db.iscritti_templates.find({"org_id": user["org_id"]}, {"_id": 0}).sort("created_at", -1).to_list(50)
        tpl = next((t for t in templates if set(t["mapping"]) == set(cols)), None)
        doc = {"id": iid, "org_id": user["org_id"], "filename": (file.filename or "")[:200], "ext": ext, "status": "caricato",
               "columns": cols, "total_rows": len(rows), "created_at": _now(), "created_by": {"user_id": user["user_id"], "name": user.get("name")}}
        await db.iscritti_imports.insert_one(dict(doc))
        return {"id": iid, "filename": doc["filename"], "columns": cols, "sample": rows[:5], "total_rows": len(rows),
                "suggested": tpl["mapping"] if tpl else suggest(cols, st.get("custom_fields", [])), "template": tpl and tpl["name"], "templates": templates}

    async def _import(org_id: str, iid: str) -> dict:
        imp = await db.iscritti_imports.find_one({"id": iid, "org_id": org_id}, {"_id": 0})
        if not imp:
            raise HTTPException(status_code=404, detail="Importazione non trovata")
        return imp

    async def _file(iid: str) -> bytes:
        chunks = await db.iscritti_import_files.find({"import_id": iid}, {"_id": 0, "data": 1}).sort("n", 1).to_list(20)
        if not chunks:
            raise HTTPException(status_code=410, detail="File non più disponibile: ricaricalo")
        return b"".join(c["data"] for c in chunks)

    async def _analyze(org_id: str, imp: dict, mapping: dict, edizione_id: str, gara: str):
        cols, rows = await asyncio.to_thread(read_rows, await _file(imp["id"]), imp["ext"])
        seen, stats, errors, missing = {}, {"total": len(rows), "new": 0, "existing": 0, "duplicates": 0, "errors": 0, "already_in_edition": 0}, [], {}
        parsed = []
        for i, row in enumerate(rows, start=2):
            p, part, custom, errs = parse_row(cols, row, mapping)
            for f in ("nome", "cognome", "email"):
                if not p.get(f):
                    missing[f] = missing.get(f, 0) + 1
            if errs and not _dedup_key(p):
                stats["errors"] += 1
                errors.append({"row": i, "errors": errs})
                continue
            if errs:
                errors.append({"row": i, "errors": errs, "warning": True})
            key = _dedup_key(p)
            race = part.get("gara") or gara
            if (key, race) in seen:
                stats["duplicates"] += 1
                continue
            seen[(key, race)] = True
            parsed.append((i, key, p, part, custom, race))
        keys = list({k for _, k, *_ in parsed})
        existing = {}
        for j in range(0, len(keys), 1000):
            async for d in db.iscritti.find({"org_id": org_id, "dedup_key": {"$in": keys[j:j + 1000]}}, {"_id": 0, "id": 1, "dedup_key": 1}):
                existing[d["dedup_key"]] = d["id"]
        in_ed = set()
        if existing:
            async for d in db.iscritti_partecipazioni.find({"org_id": org_id, "edizione_id": edizione_id, "iscritto_id": {"$in": list(existing.values())}},
                                                          {"_id": 0, "iscritto_id": 1, "gara": 1}):
                in_ed.add((d["iscritto_id"], d.get("gara") or ""))
        for _, k, *_rest in parsed:
            race = _rest[-1]
            if k in existing:
                stats["existing"] += 1
                if (existing[k], race or "") in in_ed:
                    stats["already_in_edition"] += 1
            else:
                stats["new"] += 1
        return stats, errors, missing, parsed, existing

    @r.post("/iscritti/imports/{iid}/preview")
    async def preview(iid: str, body: PreviewIn, user: dict = Depends(editor)):
        imp = await _import(user["org_id"], iid)
        if imp["status"] not in ("caricato", "analizzato"):
            raise HTTPException(status_code=400, detail="Importazione già avviata")
        ed = await _edizione(user["org_id"], body.edizione_id)
        mapping = {c: (body.mapping.get(c) or "") for c in imp["columns"]}
        if not any(v in ("email", "nome", "cognome") for v in mapping.values()):
            raise HTTPException(status_code=400, detail="Collega almeno le colonne Email oppure Nome e Cognome")
        if body.new_fields:
            st = await settings(user["org_id"])
            cf = st.get("custom_fields", [])
            for label in body.new_fields:
                key = re.sub(r"[^a-z0-9]+", "_", _norm(label)).strip("_")[:40]
                if key and key not in {c["key"] for c in cf}:
                    cf.append({"key": key, "label": label.strip()[:60]})
                if key and label in mapping:
                    mapping[label] = f"custom:{key}"
            await db.iscritti_settings.update_one({"org_id": user["org_id"]}, {"$set": {"custom_fields": cf}}, upsert=True)
        stats, errors, missing, _, _ = await _analyze(user["org_id"], imp, mapping, ed["id"], body.gara)
        await db.iscritti_imports.update_one({"id": iid}, {"$set": {"status": "analizzato", "mapping": mapping, "edizione_id": ed["id"], "event_id": ed["event_id"],
                                                                    "anno": ed["anno"], "edizione": ed["edizione"], "gara": body.gara.strip(), "preview": stats}})
        return {"stats": stats, "errors": errors[:200], "missing": missing}

    async def _run_import(org_id: str, iid: str, mode: str):
        imp = await db.iscritti_imports.find_one({"id": iid}, {"_id": 0})
        res = {"created": 0, "updated": 0, "unchanged": 0, "participations": 0, "skipped_participations": 0, "duplicates": 0, "errors": 0}
        try:
            stats, errors, _, parsed, existing = await _analyze(org_id, imp, imp["mapping"], imp["edizione_id"], imp.get("gara") or "")
            res["duplicates"], res["errors"] = stats["duplicates"], stats["errors"]
            now = _now()
            for n, (row, key, p, part, custom, race) in enumerate(parsed, start=1):
                pid = existing.get(key)
                if pid is None:
                    pid = uuid.uuid4().hex
                    doc = {"id": pid, "org_id": org_id, "dedup_key": key, **{f: p.get(f, "") for f in PERSON_FIELDS if f not in ("consenso_marketing",)},
                           "consenso_marketing": bool(p.get("consenso_marketing")), "custom": custom, "unsub_token": secrets.token_urlsafe(24),
                           "brevo_status": "da_sincronizzare" if p.get("consenso_marketing") and p.get("email") else "escluso_consenso",
                           "consent_history": [], "created_at": now, "updated_at": now, "source_import_id": iid}
                    if doc["consenso_marketing"]:
                        doc["consenso_fonte"] = p.get("consenso_fonte") or f"Importazione file «{imp['filename']}»"
                        doc["consenso_data"] = p.get("consenso_data") or now[:10]
                        doc["consent_history"] = [{"at": now, "value": True, "source": doc["consenso_fonte"], "by": imp["created_by"].get("user_id")}]
                    try:
                        await db.iscritti.insert_one(doc)
                        existing[key] = pid
                        res["created"] += 1
                    except Exception:  # noqa: BLE001 - creato in parallelo da un'altra importazione
                        pid = (await db.iscritti.find_one({"org_id": org_id, "dedup_key": key}, {"_id": 0, "id": 1}))["id"]
                elif mode == "update":
                    cur = await db.iscritti.find_one({"id": pid}, {"_id": 0})
                    upd = {f: p[f] for f in PERSON_FIELDS if f not in ("consenso_marketing", "consenso_fonte", "consenso_data") and p.get(f) and p[f] != cur.get(f)}
                    for k, v in custom.items():
                        if v and (cur.get("custom") or {}).get(k) != v:
                            upd[f"custom.{k}"] = v
                    push = None
                    if p.get("consenso_marketing") and not cur.get("consenso_marketing") and not cur.get("unsubscribed"):
                        upd.update(consenso_marketing=True, consenso_fonte=p.get("consenso_fonte") or f"Importazione file «{imp['filename']}»",
                                   consenso_data=p.get("consenso_data") or now[:10], brevo_status="da_sincronizzare")
                        push = {"consent_history": {"at": now, "value": True, "source": upd["consenso_fonte"], "by": imp["created_by"].get("user_id")}}
                    if upd:
                        if cur.get("consenso_marketing") and cur.get("brevo_status") == "sincronizzato" and ("nome" in upd or "cognome" in upd):
                            upd["brevo_status"] = "da_sincronizzare"
                        upd["updated_at"] = now
                        await db.iscritti.update_one({"id": pid}, {"$set": upd, **({"$push": push} if push else {})})
                        res["updated"] += 1
                    else:
                        res["unchanged"] += 1
                else:
                    res["unchanged"] += 1
                pdoc = {"org_id": org_id, "iscritto_id": pid, "edizione_id": imp["edizione_id"], "event_id": imp["event_id"], "anno": imp["anno"],
                        "edizione": imp["edizione"], "gara": race or "", "tipologia": part.get("tipologia", ""), "stato": (part.get("stato") or "confermata")[:40],
                        "data_iscrizione": part.get("data_iscrizione"), "import_id": iid}
                r0 = await db.iscritti_partecipazioni.update_one({"org_id": org_id, "iscritto_id": pid, "edizione_id": imp["edizione_id"], "gara": race or ""},
                                                                 {"$setOnInsert": {"id": uuid.uuid4().hex, "created_at": now, **pdoc}}, upsert=True)
                if r0.upserted_id:
                    res["participations"] += 1
                    await db.iscritti.update_one({"id": pid, "$or": [{"last_year": {"$lt": imp["anno"]}}, {"last_year": None}]}, {"$set": {"last_year": imp["anno"]}})
                else:
                    res["skipped_participations"] += 1
                if n % 250 == 0:
                    await db.iscritti_imports.update_one({"id": iid}, {"$set": {"progress": n, "result": res}})
            await db.iscritti_imports.update_one({"id": iid}, {"$set": {"status": "completato", "progress": len(parsed), "result": res,
                                                                        "errors": errors[:500], "finished_at": _now()}})
            asyncio.create_task(sync_org(org_id))
        except Exception as e:  # noqa: BLE001
            log.error("iscritti import %s failed: %s", iid, type(e).__name__)
            await db.iscritti_imports.update_one({"id": iid}, {"$set": {"status": "errore", "result": res, "error": f"Errore durante l'importazione ({type(e).__name__})", "finished_at": _now()}})
        finally:
            await db.iscritti_import_files.delete_many({"import_id": iid})

    @r.post("/iscritti/imports/{iid}/run")
    async def run(iid: str, body: RunIn, user: dict = Depends(editor)):
        imp = await _import(user["org_id"], iid)
        if imp["status"] != "analizzato":
            raise HTTPException(status_code=400, detail="Completa l'anteprima prima di importare")
        res = await db.iscritti_imports.update_one({"id": iid, "status": "analizzato"}, {"$set": {"status": "in_corso", "mode": body.mode, "started_at": _now(), "progress": 0}})
        if not res.modified_count:
            raise HTTPException(status_code=409, detail="Importazione già avviata")
        if body.template_name and body.template_name.strip():
            await db.iscritti_templates.update_one({"org_id": user["org_id"], "name": body.template_name.strip()},
                                                   {"$set": {"mapping": imp["mapping"], "updated_at": _now()}, "$setOnInsert": {"id": uuid.uuid4().hex, "created_at": _now()}}, upsert=True)
        await record_audit(user, "iscritti_import_started", org_id=user["org_id"], detail=f"File «{imp['filename']}» · {imp['total_rows']} righe · {imp.get('anno')} {imp.get('edizione')}")
        asyncio.create_task(_run_import(user["org_id"], iid, body.mode))
        return {"ok": True, "status": "in_corso"}

    @r.get("/iscritti/imports")
    async def list_imports(user: dict = Depends(member)):
        rows = await db.iscritti_imports.find({"org_id": user["org_id"], "status": {"$in": ["in_corso", "completato", "errore"]}}, {"_id": 0, "errors": 0, "columns": 0}).sort("created_at", -1).to_list(100)
        return {"imports": rows}

    @r.get("/iscritti/imports/{iid}")
    async def get_import(iid: str, user: dict = Depends(member)):
        imp = await _import(user["org_id"], iid)
        return imp

    @r.delete("/iscritti/templates/{tid}")
    async def delete_template(tid: str, user: dict = Depends(editor)):
        await db.iscritti_templates.delete_one({"id": tid, "org_id": user["org_id"]})
        return {"ok": True}

    # ---------------- consultazione ----------------
    def _filters(**kw) -> dict:
        return {k: v for k, v in kw.items() if v not in (None, "")}

    @r.get("/iscritti")
    async def list_iscritti(user: dict = Depends(member), q: str = "", event_id: str = "", anno: str = "", edizione_id: str = "", gara: str = "", stato: str = "",
                            consenso: str = "", sort: str = "cognome", dir: int = 1, page: int = Query(1, ge=1), limit: int = Query(50, ge=1, le=200)):
        f = _filters(q=q, event_id=event_id, anno=anno, edizione_id=edizione_id, gara=gara, stato=stato, consenso=consenso)
        query = await person_query(user["org_id"], f)
        total = await db.iscritti.count_documents(query)
        sort_f = sort if sort in ("cognome", "nome", "email", "societa", "created_at", "last_year") else "cognome"
        rows = await db.iscritti.find(query, {"_id": 0, "unsub_token": 0, "dedup_key": 0, "consent_history": 0}).sort([(sort_f, 1 if dir >= 0 else -1), ("nome", 1)]) \
            .skip((page - 1) * limit).to_list(limit)
        counts = {c["_id"]: c["n"] for c in await db.iscritti_partecipazioni.aggregate([{"$match": {"org_id": user["org_id"], "iscritto_id": {"$in": [x["id"] for x in rows]}}},
                                                                                         {"$group": {"_id": "$iscritto_id", "n": {"$sum": 1}}}]).to_list(500)}
        return {"items": [{**x, "partecipazioni": counts.get(x["id"], 0)} for x in rows], "total": total, "page": page, "limit": limit}

    @r.get("/iscritti/kpi")
    async def kpi(user: dict = Depends(member), event_id: str = "", anno: str = "", edizione_id: str = "", gara: str = "", stato: str = ""):
        oid = user["org_id"]
        f = _filters(event_id=event_id, anno=anno, edizione_id=edizione_id, gara=gara, stato=stato)
        pq = await participation_query(oid, f) or {"org_id": oid}
        total = await db.iscritti.count_documents(await person_query(oid, f))
        per_person = await db.iscritti_partecipazioni.aggregate([{"$match": {"org_id": oid}}, {"$group": {"_id": "$iscritto_id", "n": {"$sum": 1}, "eds": {"$addToSet": "$edizione_id"}}}]).to_list(200000)
        sel = set(await db.iscritti_partecipazioni.distinct("iscritto_id", pq))
        recurring = sum(1 for p in per_person if p["_id"] in sel and len(p["eds"]) > 1)
        by_event = await db.iscritti_partecipazioni.aggregate([{"$match": pq}, {"$group": {"_id": "$event_id", "n": {"$addToSet": "$iscritto_id"}}},
                                                                {"$project": {"n": {"$size": "$n"}}}, {"$sort": {"n": -1}}]).to_list(50)
        by_ed = await db.iscritti_partecipazioni.aggregate([{"$match": pq}, {"$group": {"_id": "$edizione_id", "n": {"$addToSet": "$iscritto_id"}}},
                                                             {"$project": {"n": {"$size": "$n"}}}, {"$sort": {"n": -1}}]).to_list(50)
        evs = {e["id"]: e.get("nome") for e in await db.events.find({"org_id": oid}, {"_id": 0, "id": 1, "nome": 1}).to_list(5000)}
        eds = {e["id"]: e for e in await db.iscritti_edizioni.find({"org_id": oid}, {"_id": 0}).to_list(2000)}
        new_since = (datetime.now(timezone.utc) - timedelta(days=30)).isoformat()
        return {"total": total, "recurring": recurring, "new_30d": await db.iscritti.count_documents({**(await person_query(oid, f)), "created_at": {"$gte": new_since}}),
                "with_consent": await db.iscritti.count_documents({**(await person_query(oid, f)), "consenso_marketing": True}),
                "by_event": [{"event_id": x["_id"], "nome": evs.get(x["_id"], "—"), "n": x["n"]} for x in by_event],
                "by_edition": [{"edizione_id": x["_id"], "label": f"{(eds.get(x['_id']) or {}).get('event_nome', '—')} {(eds.get(x['_id']) or {}).get('anno', '')} {(eds.get(x['_id']) or {}).get('edizione', '')}".strip(), "n": x["n"]} for x in by_ed]}

    @r.get("/iscritti/export")
    async def export(user: dict = Depends(editor), format: str = "csv", q: str = "", event_id: str = "", anno: str = "", edizione_id: str = "", gara: str = "", stato: str = "", consenso: str = ""):
        f = _filters(q=q, event_id=event_id, anno=anno, edizione_id=edizione_id, gara=gara, stato=stato, consenso=consenso)
        rows = await db.iscritti.find(await person_query(user["org_id"], f), {"_id": 0}).sort([("cognome", 1), ("nome", 1)]).to_list(200000)
        cf = (await settings(user["org_id"])).get("custom_fields", [])
        parts = {}
        async for p in db.iscritti_partecipazioni.find({"org_id": user["org_id"], "iscritto_id": {"$in": [x["id"] for x in rows]}}, {"_id": 0}):
            parts.setdefault(p["iscritto_id"], []).append(f"{p.get('anno')} {p.get('edizione') or ''} {p.get('gara') or ''}".strip())
        head = [LABELS[k] for k in PERSON_FIELDS] + [c["label"] for c in cf] + ["Partecipazioni"]
        data = [[("Sì" if x.get(k) else "No") if k == "consenso_marketing" else (x.get(k) or "") for k in PERSON_FIELDS]
                + [(x.get("custom") or {}).get(c["key"], "") for c in cf] + [" | ".join(parts.get(x["id"], []))] for x in rows]
        await record_audit(user, "iscritti_export", org_id=user["org_id"], detail=f"Export {format.upper()} di {len(rows)} iscritti")
        if format == "xlsx":
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Iscritti"
            ws.append(head)
            for d in data:
                ws.append([str(v) for v in d])
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                     headers={"Content-Disposition": 'attachment; filename="iscritti.xlsx"'})
        out = io.StringIO()
        w = csv.writer(out, delimiter=";")
        w.writerow(head)
        w.writerows(data)
        return StreamingResponse(iter(["\ufeff" + out.getvalue()]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="iscritti.csv"'})

    @r.get("/iscritti/{pid}")
    async def get_iscritto(pid: str, user: dict = Depends(member)):
        p = await db.iscritti.find_one({"id": pid, "org_id": user["org_id"]}, {"_id": 0})
        if not p:
            raise HTTPException(status_code=404, detail="Iscritto non trovato")
        parts = await db.iscritti_partecipazioni.find({"org_id": user["org_id"], "iscritto_id": pid}, {"_id": 0}).sort("anno", -1).to_list(500)
        eds = {e["id"]: e for e in await db.iscritti_edizioni.find({"org_id": user["org_id"], "id": {"$in": [x["edizione_id"] for x in parts]}}, {"_id": 0}).to_list(500)}
        return {**_pub(p), "partecipazioni": [{**x, "event_nome": (eds.get(x["edizione_id"]) or {}).get("event_nome")} for x in parts]}

    def _clean(body: IscrittoIn) -> dict:
        d = body.model_dump(exclude={"partecipazione", "custom", "consenso_marketing", "consenso_fonte"})
        d = {k: (v or "").strip() for k, v in d.items()}
        d["email"] = d["email"].lower()
        d["codice_fiscale"] = d["codice_fiscale"].upper().replace(" ", "")
        if d["email"] and not EMAIL_RX.match(d["email"]):
            raise HTTPException(status_code=400, detail="Email non valida")
        if d["data_nascita"] and not _date(d["data_nascita"]):
            raise HTTPException(status_code=400, detail="Data di nascita non valida")
        if not (d["email"] or (d["nome"] and d["cognome"])):
            raise HTTPException(status_code=400, detail="Inserisci almeno l'email oppure nome e cognome")
        return d

    async def _add_part(org_id: str, pid: str, part: dict, now: str):
        ed = await _edizione(org_id, part.get("edizione_id", ""))
        await db.iscritti_partecipazioni.update_one({"org_id": org_id, "iscritto_id": pid, "edizione_id": ed["id"], "gara": part.get("gara", "")}, {"$setOnInsert": {
            "id": uuid.uuid4().hex, "org_id": org_id, "iscritto_id": pid, "edizione_id": ed["id"], "event_id": ed["event_id"], "anno": ed["anno"],
            "edizione": ed["edizione"], "gara": part.get("gara", ""), "tipologia": part.get("tipologia", ""), "stato": part.get("stato") or "confermata",
            "data_iscrizione": _date(part.get("data_iscrizione", "")) or now[:10], "created_at": now}}, upsert=True)
        await db.iscritti.update_one({"id": pid, "$or": [{"last_year": {"$lt": ed["anno"]}}, {"last_year": None}]}, {"$set": {"last_year": ed["anno"]}})

    @r.post("/iscritti")
    async def create_iscritto(body: IscrittoIn, user: dict = Depends(editor)):
        d, now = _clean(body), _now()
        key = _dedup_key(d)
        if not key:
            raise HTTPException(status_code=400, detail="Per evitare duplicati servono email, codice fiscale o nome+cognome+data di nascita")
        if await db.iscritti.find_one({"org_id": user["org_id"], "dedup_key": key}, {"_id": 1}):
            raise HTTPException(status_code=409, detail="Questa persona è già presente nell'archivio")
        doc = {"id": uuid.uuid4().hex, "org_id": user["org_id"], "dedup_key": key, **d, "custom": {k: v[:500] for k, v in body.custom.items()},
               "consenso_marketing": body.consenso_marketing, "consenso_fonte": body.consenso_fonte.strip() if body.consenso_marketing else "",
               "consenso_data": now[:10] if body.consenso_marketing else "", "unsub_token": secrets.token_urlsafe(24),
               "consent_history": [{"at": now, "value": True, "source": body.consenso_fonte or "Inserimento manuale", "by": user["user_id"]}] if body.consenso_marketing else [],
               "brevo_status": "da_sincronizzare" if body.consenso_marketing and d["email"] else "escluso_consenso", "created_at": now, "updated_at": now}
        await db.iscritti.insert_one(dict(doc))
        if body.partecipazione and body.partecipazione.get("edizione_id"):
            await _add_part(user["org_id"], doc["id"], body.partecipazione, now)
        asyncio.create_task(sync_org(user["org_id"]))
        return _pub(doc)

    @r.put("/iscritti/{pid}")
    async def update_iscritto(pid: str, body: IscrittoIn, user: dict = Depends(editor)):
        cur = await db.iscritti.find_one({"id": pid, "org_id": user["org_id"], "anonymized": {"$ne": True}}, {"_id": 0})
        if not cur:
            raise HTTPException(status_code=404, detail="Iscritto non trovato")
        d, now = _clean(body), _now()
        key = _dedup_key(d) or cur["dedup_key"]
        if key != cur["dedup_key"] and await db.iscritti.find_one({"org_id": user["org_id"], "dedup_key": key}, {"_id": 1}):
            raise HTTPException(status_code=409, detail="Esiste già un altro iscritto con questi dati")
        upd = {**d, "dedup_key": key, "custom": {k: v[:500] for k, v in body.custom.items()}, "updated_at": now}
        push = None
        if body.consenso_marketing != bool(cur.get("consenso_marketing")):
            if body.consenso_marketing and cur.get("unsubscribed"):
                raise HTTPException(status_code=400, detail="L'iscritto si è disiscritto tramite il link: il consenso può essere riattivato solo dall'interessato")
            src = body.consenso_fonte.strip() or ("Modifica manuale" if body.consenso_marketing else "Revoca manuale")
            upd.update(consenso_marketing=body.consenso_marketing, consenso_fonte=src, consenso_data=now[:10])
            push = {"consent_history": {"at": now, "value": body.consenso_marketing, "source": src, "by": user["user_id"]}}
            if body.consenso_marketing:
                pass
        upd["brevo_status"] = "da_sincronizzare" if (upd.get("consenso_marketing", cur.get("consenso_marketing")) and d["email"]) else (
            "da_rimuovere" if cur.get("brevo_status") == "sincronizzato" else "escluso_consenso")
        await db.iscritti.update_one({"id": pid}, {"$set": upd, **({"$push": push} if push else {})})
        if body.partecipazione and body.partecipazione.get("edizione_id"):
            await _add_part(user["org_id"], pid, body.partecipazione, now)
        asyncio.create_task(sync_org(user["org_id"]))
        return await get_iscritto(pid, user)

    @r.delete("/iscritti/partecipazioni/{part_id}")
    async def delete_part(part_id: str, user: dict = Depends(editor)):
        await db.iscritti_partecipazioni.delete_one({"id": part_id, "org_id": user["org_id"]})
        return {"ok": True}

    async def _anonymize(org_id: str, pid: str, reason: str):
        p = await db.iscritti.find_one({"id": pid, "org_id": org_id}, {"_id": 0, "email": 1, "brevo_status": 1})
        if not p:
            return False
        if p.get("brevo_status") == "sincronizzato" and p.get("email"):
            await _brevo_remove(org_id, p["email"])
        await db.iscritti.update_one({"id": pid}, {"$set": {"anonymized": True, "anonymized_at": _now(), "anonymized_reason": reason, "nome": "", "cognome": "",
                                                            "email": "", "cellulare": "", "data_nascita": "", "codice_fiscale": "", "numero_tessera": "",
                                                            "note": "", "custom": {}, "dedup_key": f"anon:{pid}", "consenso_marketing": False,
                                                            "consent_history": [], "brevo_status": "anonimizzato", "unsub_token": None}})
        return True

    @r.post("/iscritti/{pid}/anonymize")
    async def anonymize(pid: str, user: dict = Depends(editor)):
        if not await _anonymize(user["org_id"], pid, "richiesta"):
            raise HTTPException(status_code=404, detail="Iscritto non trovato")
        await record_audit(user, "iscritto_anonymized", org_id=user["org_id"], detail=f"Iscritto {pid} anonimizzato su richiesta")
        return {"ok": True}

    @r.delete("/iscritti/{pid}")
    async def delete_iscritto(pid: str, user: dict = Depends(editor)):
        p = await db.iscritti.find_one({"id": pid, "org_id": user["org_id"]}, {"_id": 0, "email": 1, "brevo_status": 1})
        if not p:
            raise HTTPException(status_code=404, detail="Iscritto non trovato")
        if p.get("brevo_status") == "sincronizzato" and p.get("email"):
            await _brevo_remove(user["org_id"], p["email"])
        await db.iscritti_partecipazioni.delete_many({"org_id": user["org_id"], "iscritto_id": pid})
        await db.iscritti.delete_one({"id": pid})
        await record_audit(user, "iscritto_deleted", org_id=user["org_id"], detail=f"Iscritto {pid} cancellato con le partecipazioni")
        return {"ok": True}

    # ---------------- Brevo ----------------
    async def _ensure_list(c, org: dict, st: dict) -> int:
        if st.get("brevo_list_id"):
            return st["brevo_list_id"]
        name = f"CRMEvent – {org.get('nome') or 'Organizzazione'} – Iscritti"
        found = next((l for l in await c.lists() if (l.get("name") or "").strip() == name), None)
        lid = found["id"] if found else (await c.create_list(name, await brevo_org_lists._folder_id(c)))["id"]
        await db.iscritti_settings.update_one({"org_id": org["id"]}, {"$set": {"brevo_list_id": lid, "brevo_list_name": name}}, upsert=True)
        return lid

    async def _brevo_remove(org_id: str, email: str):
        st = await settings(org_id)
        if not (st.get("brevo_list_id") and brevo_client.is_configured()):
            return
        try:
            async with brevo_client.BrevoClient() as c:
                await c._req("POST", f"/contacts/lists/{st['brevo_list_id']}/contacts/remove", json={"emails": [email]})
        except Exception as e:  # noqa: BLE001
            log.warning("iscritti brevo remove failed: %s", type(e).__name__)

    async def sync_org(org_id: str, force_ids: Optional[list] = None) -> dict:
        """Sincronizza i contatti con consenso (o quelli forzati dal Super Admin). Mai riattiva disiscritti o blacklist."""
        org = await db.organizations.find_one({"id": org_id}, {"_id": 0, "id": 1, "nome": 1})
        if not org:
            return {}
        now = _now()
        if not brevo_client.is_configured():
            out = {"status": "non_configurato", "at": now, "error": "Brevo non configurato"}
            await db.iscritti_settings.update_one({"org_id": org_id}, {"$set": {"brevo_last_sync": out}}, upsert=True)
            return out
        q = {"org_id": org_id, "anonymized": {"$ne": True}, "email": {"$nin": [None, ""]}}
        q.update({"id": {"$in": force_ids}} if force_ids else {"brevo_status": {"$in": ["da_sincronizzare", "errore", "rate_limit", "da_rimuovere"]}})
        counts = {"sincronizzati": 0, "disiscritti": 0, "rimossi": 0, "errori": 0}
        status, error = "ok", None
        try:
            async with brevo_client.BrevoClient() as c:
                lid = await _ensure_list(c, org, await settings(org_id))
                async for p in db.iscritti.find(q, {"_id": 0}):
                    try:
                        if p.get("brevo_status") == "da_rimuovere" and not force_ids:
                            await c._req("POST", f"/contacts/lists/{lid}/contacts/remove", json={"emails": [p["email"]]})
                            await db.iscritti.update_one({"id": p["id"]}, {"$set": {"brevo_status": "escluso_consenso", "brevo_synced_at": now}})
                            counts["rimossi"] += 1
                            continue
                        if not force_ids and not p.get("consenso_marketing"):
                            continue
                        contact = await c.get_contact(p["email"])
                        if contact and (contact.get("emailBlacklisted") or lid in (contact.get("listUnsubscribed") or [])):
                            await db.iscritti.update_one({"id": p["id"]}, {"$set": {"brevo_status": "disiscritto", "brevo_synced_at": now}})
                            counts["disiscritti"] += 1
                            continue
                        attrs = {k: v for k, v in {"NOME": p.get("nome"), "COGNOME": p.get("cognome")}.items() if v}
                        if contact:
                            await c.update_contact(p["email"], lid, attrs or None)
                        else:
                            await c.create_contact(p["email"], lid, attrs or None)
                        await db.iscritti.update_one({"id": p["id"]}, {"$set": {"brevo_status": "sincronizzato", "brevo_synced_at": now, "brevo_error": None,
                                                                                **({"brevo_forced": True} if force_ids else {})}})
                        counts["sincronizzati"] += 1
                    except brevo_client.BrevoError as e:
                        if e.status == 429:
                            await db.iscritti.update_one({"id": p["id"]}, {"$set": {"brevo_status": "rate_limit"}})
                            status, error = "rate_limit", "Limite API Brevo raggiunto: nuovo tentativo automatico"
                            break
                        await db.iscritti.update_one({"id": p["id"]}, {"$set": {"brevo_status": "errore", "brevo_error": f"HTTP {e.status}"}})
                        counts["errori"] += 1
                    await asyncio.sleep(0.05)
        except Exception as e:  # noqa: BLE001
            status, error = "errore", f"Sincronizzazione non riuscita ({type(e).__name__})"
            log.warning("iscritti sync %s failed: %s", org_id, type(e).__name__)
        if status == "ok" and counts["errori"]:
            status, error = "parziale", f"{counts['errori']} contatti non sincronizzati"
        out = {"status": status, "at": now, "error": error, "counts": counts}
        await db.iscritti_settings.update_one({"org_id": org_id}, {"$set": {"brevo_last_sync": out}}, upsert=True)
        return out

    @r.post("/iscritti/brevo/sync")
    async def sync_now(user: dict = Depends(editor)):
        return await sync_org(user["org_id"])

    @r.post("/iscritti/brevo/segment")
    async def segment(body: SegmentIn, user: dict = Depends(editor)):
        """Lista Brevo dedicata a un evento/edizione (solo iscritti già sincronizzati con consenso)."""
        f = {k: v for k, v in body.filters.items() if k in ("event_id", "anno", "edizione_id", "gara", "stato") and v}
        if not f:
            raise HTTPException(status_code=400, detail="Seleziona almeno evento, anno o edizione")
        if not brevo_client.is_configured():
            raise HTTPException(status_code=400, detail="Brevo non configurato")
        org = await db.organizations.find_one({"id": user["org_id"]}, {"_id": 0, "nome": 1})
        ed = await db.iscritti_edizioni.find_one({"id": f.get("edizione_id"), "org_id": user["org_id"]}, {"_id": 0}) if f.get("edizione_id") else None
        ev = await db.events.find_one({"id": f.get("event_id") or (ed or {}).get("event_id"), "org_id": user["org_id"]}, {"_id": 0, "nome": 1}) or {}
        label = " ".join(x for x in [ev.get("nome"), str((ed or {}).get("anno") or f.get("anno") or ""), (ed or {}).get("edizione"), f.get("gara")] if x)
        name = f"CRMEvent – {org.get('nome')} – Iscritti – {label or 'Selezione'}"
        emails = [p["email"] for p in await select_recipients(user["org_id"], f) if p.get("email")]
        synced = set(await db.iscritti.distinct("email", {"org_id": user["org_id"], "email": {"$in": emails}, "brevo_status": "sincronizzato"}))
        emails = [e for e in emails if e in synced]
        try:
            async with brevo_client.BrevoClient() as c:
                found = next((l for l in await c.lists() if (l.get("name") or "").strip() == name), None)
                lid = found["id"] if found else (await c.create_list(name, await brevo_org_lists._folder_id(c)))["id"]
                for i in range(0, len(emails), 150):
                    await c._req("POST", f"/contacts/lists/{lid}/contacts/add", json={"emails": emails[i:i + 150]})
        except brevo_client.BrevoError as e:
            raise HTTPException(status_code=400, detail=f"Brevo: errore {e.status}")
        await db.iscritti_segments.update_one({"org_id": user["org_id"], "list_id": lid}, {"$set": {"name": name, "filters": f, "count": len(emails), "updated_at": _now()}}, upsert=True)
        return {"list_id": lid, "name": name, "count": len(emails)}

    @r.get("/public/iscritti/unsubscribe/{token}", response_class=HTMLResponse)
    async def public_unsub(token: str):
        p = await db.iscritti.find_one({"unsub_token": token}, {"_id": 0}) if token and len(token) > 20 else None
        if p and p.get("consenso_marketing"):
            now = _now()
            await db.iscritti.update_one({"id": p["id"]}, {"$set": {"consenso_marketing": False, "unsubscribed": True, "consenso_data": now[:10], "consenso_fonte": "Disiscrizione via link",
                                                                    "brevo_status": "da_rimuovere" if p.get("brevo_status") == "sincronizzato" else "escluso_consenso"},
                                                           "$push": {"consent_history": {"at": now, "value": False, "source": "Disiscrizione via link"}}})
            asyncio.create_task(sync_org(p["org_id"]))
        return HTMLResponse("<html><body style='font-family:sans-serif;padding:40px;text-align:center'><h2>Disiscrizione completata</h2>"
                            "<p>Non riceverai più comunicazioni promozionali da questa organizzazione.</p></body></html>")

    # ---------------- Super Admin ----------------
    @r.get("/platform/iscritti/orgs")
    async def sa_orgs(admin: dict = Depends(sa)):
        svc = await db.marketplace_services.find_one({"key": SERVICE_KEY}, {"_id": 0, "id": 1, "price": 1, "price_type": 1, "status": 1})
        act = {s["org_id"]: s for s in await db.org_services.find({"service": SERVICE_KEY}, {"_id": 0}).to_list(5000)}
        bought = {p["org_id"]: p for p in await db.marketplace_purchases.find({"service_id": (svc or {}).get("id"), "status": {"$in": ["active", "past_due"]}}, {"_id": 0}).to_list(5000)} if svc else {}
        out = []
        for o in await db.organizations.find({}, {"_id": 0, "id": 1, "nome": 1}).sort("nome", 1).to_list(5000):
            a, b = act.get(o["id"]) or {}, bought.get(o["id"])
            if not (a or b):
                continue
            st = await settings(o["id"])
            last = await db.iscritti_imports.find_one({"org_id": o["id"], "status": {"$in": ["in_corso", "completato", "errore"]}}, {"_id": 0, "filename": 1, "status": 1, "created_at": 1}, sort=[("created_at", -1)])
            out.append({"org_id": o["id"], "nome": o["nome"], "active": bool(a.get("active")) or bool(b), "via": "acquisto" if b else "super_admin",
                        "iscritti": await db.iscritti.count_documents({"org_id": o["id"], "anonymized": {"$ne": True}}),
                        "brevo_errors": await db.iscritti.count_documents({"org_id": o["id"], "brevo_status": {"$in": ["errore", "rate_limit"]}}),
                        "not_consented": await db.iscritti.count_documents({"org_id": o["id"], "brevo_status": "escluso_consenso", "email": {"$nin": [None, ""]}, "unsubscribed": {"$ne": True}}),
                        "last_import": last, "brevo": {k: st.get(k) for k in ("brevo_list_id", "brevo_list_name", "brevo_last_sync")}})
        return {"service": svc, "orgs": out}

    @r.post("/platform/iscritti/orgs/{org_id}/activate")
    async def sa_activate(org_id: str, body: ActivateIn, admin: dict = Depends(sa)):
        org = await db.organizations.find_one({"id": org_id}, {"_id": 0, "nome": 1})
        if not org:
            raise HTTPException(status_code=404, detail="Organizzazione non trovata")
        await db.org_services.update_one({"org_id": org_id, "service": SERVICE_KEY}, {"$set": {
            "active": body.active, "updated_at": _now(), "billing": "attivazione_super_admin",
            **({"activated_at": _now(), "activated_by": admin.get("email")} if body.active else {"deactivated_at": _now()})}}, upsert=True)
        if body.active:
            asyncio.create_task(sync_org(org_id))
        await record_audit(admin, "iscritti_service_" + ("activated" if body.active else "deactivated"), org_id=org_id, org_name=org.get("nome"))
        return {"ok": True}

    @r.post("/platform/iscritti/orgs/{org_id}/sync")
    async def sa_sync(org_id: str, admin: dict = Depends(sa)):
        await db.iscritti.update_many({"org_id": org_id, "brevo_status": {"$in": ["errore", "rate_limit"]}}, {"$set": {"brevo_status": "da_sincronizzare"}})
        res = await sync_org(org_id)
        await record_audit(admin, "iscritti_brevo_resync", org_id=org_id, detail=str(res.get("counts")))
        return res

    @r.post("/platform/iscritti/orgs/{org_id}/force-sync")
    async def sa_force(org_id: str, admin: dict = Depends(sa)):
        """Sincronizza su Brevo anche gli iscritti senza consenso registrato (mai i disiscritti): azione tracciata."""
        ids = await db.iscritti.distinct("id", {"org_id": org_id, "brevo_status": "escluso_consenso", "email": {"$nin": [None, ""]},
                                                "unsubscribed": {"$ne": True}, "anonymized": {"$ne": True}})
        res = await sync_org(org_id, force_ids=ids) if ids else {"status": "ok", "counts": {"sincronizzati": 0}}
        await record_audit(admin, "iscritti_brevo_force_sync", org_id=org_id, detail=f"{len(ids)} iscritti senza consenso inviati a Brevo dal Super Admin")
        return res

    # ---------------- cron giornaliero: retry sync + conservazione ----------------
    @r.post("/cron/iscritti-tick")
    async def cron(authorization: str = Header(default="")):
        # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
        expected = deps["cron_secret"]
        if not expected or not secrets.compare_digest(authorization or "", f"Bearer {expected}"):
            raise HTTPException(status_code=401, detail="unauthorized")

        async def _run():
            for oid in await db.iscritti.distinct("org_id", {"brevo_status": {"$in": ["da_sincronizzare", "errore", "rate_limit", "da_rimuovere"]}}):
                if await is_active(oid):
                    await sync_org(oid)
            limit_year = datetime.now(timezone.utc).year - RETENTION_YEARS
            grace = (datetime.now(timezone.utc) - timedelta(days=90)).isoformat()  # storico appena importato: 90 giorni per verificarlo
            await db.iscritti_import_files.delete_many({"created_at": {"$lt": (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()}})
            async for p in db.iscritti.find({"anonymized": {"$ne": True}, "last_year": {"$lt": limit_year}, "created_at": {"$lt": grace}}, {"_id": 0, "id": 1, "org_id": 1}):
                await _anonymize(p["org_id"], p["id"], f"conservazione: nessuna partecipazione da oltre {RETENTION_YEARS} anni")
        asyncio.create_task(_run())
        return {"accepted": True}

    async def ensure_indexes():
        await db.iscritti.create_index([("org_id", 1), ("dedup_key", 1)], unique=True)
        await db.iscritti.create_index("id", unique=True)
        await db.iscritti.create_index("unsub_token", sparse=True)
        await db.iscritti_partecipazioni.create_index([("org_id", 1), ("iscritto_id", 1), ("edizione_id", 1), ("gara", 1)], unique=True)
        await db.iscritti_partecipazioni.create_index([("org_id", 1), ("edizione_id", 1)])
        await db.iscritti_edizioni.create_index([("org_id", 1), ("event_id", 1)])
        await db.iscritti_imports.create_index([("org_id", 1), ("created_at", -1)])
        await db.iscritti_import_files.create_index([("import_id", 1), ("n", 1)])
        if not await db.marketplace_meta.find_one({"id": "iscritti_v1"}):
            await db.marketplace_services.update_one({"key": SERVICE_KEY}, {"$set": {
                "name": "Elenco Iscritti", "status": "available", "tech_ready": True, "purchasable": True, "visible": True, "price_type": "monthly",
                "description": "Importa gli iscritti da Excel e CSV, crea l'archivio storico dei partecipanti per evento, anno ed edizione e sincronizza i contatti con consenso su Brevo.",
                "terms": "Importazione guidata XLSX/XLS/CSV · archivio unico senza duplicati · storico partecipazioni per edizione · filtri, KPI ed export · lista Brevo dedicata · dati disponibili agli altri moduli Marketplace.",
                "updated_at": _now()}})
            await db.marketplace_meta.insert_one({"id": "iscritti_v1", "applied_at": _now()})

    return {"router": r, "ensure_indexes": ensure_indexes, "is_active": is_active, "select_recipients": select_recipients,
            "unsub_url": unsub_url, "sync_org": sync_org}
